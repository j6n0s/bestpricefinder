"""Parser for GVH's official daily bulk file ``arfigyelo_napi_termekadatok.xlsx`` (BUILD-BRIEF §4.1).

One row per product × chain (about 25,400 rows, all categories). Traps handled here (trap 4):

* **Stale copies.** The only sheet is named ``Napi feltöltések - YYYY-MM-DD``; that date must equal the
  expected data date, else :class:`StaleXlsxError`. (The cache-busting query string is the client's job.)
* **Numbers are text with a decimal comma** (``"1699,0000"``). They are parsed to :class:`Decimal` exactly;
  a ``.``, a thousands separator or any other character is rejected rather than guessed. Real numeric cells
  are accepted too.
* **Header validation.** Columns are located by their exact (NFC, whitespace-collapsed) header text. A
  missing or duplicated documented header raises :class:`GvhSchemaError`; unknown extra headers and a
  changed column order are reported as drift (the mapping by name stays correct).

Prices are the "daily price" (DISCOUNTED included, LOYALTY excluded); min/max are across the chain's stores,
so the minimum can be a single-store outlier (trap 1). There are no per-store prices.
"""

from __future__ import annotations

import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.xml import DEFUSEDXML

from bpf.gvh.errors import GvhError, GvhSchemaError, StaleXlsxError

SHEET_PREFIX = "Napi feltöltések - "
_SHEET_RE = re.compile(r"Napi feltöltések - (\d{4}-\d{2}-\d{2})")

#: The documented columns, in the documented order (research/gvh-location/notes.md §5, verify.md §2).
HEADERS: tuple[str, ...] = (
    "Termék azonosító",
    "Termék név",
    "Kategória azonosító",
    "Kategória név",
    "Üzletlánc név",
    "Egység",
    "Kiszerelés",
    "Minimum ár",
    "Maximum ár",
    "Minimum egységár",
    "Maximum egységár",
    "Hány üzletláncban elérhető",
    "Hány boltban elérhető",
    "Üzletlánc összes boltja",
)

_DECIMAL_TEXT = re.compile(r"[+-]?\d+(?:,\d+)?")
_MAX_CELL_ECHO = 30


@dataclass(frozen=True, slots=True)
class XlsxRow:
    """One product × chain row. Prices in HUF; ``*_unit_price`` per ``unit``; ``packaging`` in ``unit``."""

    product_id: str
    product_name: str
    category_id: int
    category_name: str | None
    chain_name: str
    unit: str
    packaging: Decimal | None
    min_price: Decimal | None
    max_price: Decimal | None
    min_unit_price: Decimal | None
    max_unit_price: Decimal | None
    chain_count: int | None
    stores_carrying: int | None
    chain_store_count: int | None


@dataclass(frozen=True, slots=True)
class DailyXlsx:
    data_date: date
    sheet_name: str
    rows: tuple[XlsxRow, ...]
    drift: tuple[str, ...] = ()


def _norm(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", text).split())


def sheet_date(sheet_name: str) -> date | None:
    """The date in a ``Napi feltöltések - YYYY-MM-DD`` sheet name, or None if the name does not match."""
    match = _SHEET_RE.fullmatch(_norm(sheet_name))
    if match is None:
        return None
    try:
        return date.fromisoformat(match.group(1))
    except ValueError:
        return None


def parse_decimal_cell(value: object) -> Decimal | None:
    """``"1699,0000"`` -> Decimal('1699.0000'); ``""``/None -> None. Raises ValueError for anything else."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("boolean where a number was expected")
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        dec = Decimal(repr(value))
        if not dec.is_finite():
            raise ValueError("non-finite number")
        return dec
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if not _DECIMAL_TEXT.fullmatch(text):
            raise ValueError("not a decimal-comma number")
        try:
            return Decimal(text.replace(",", "."))
        except InvalidOperation as exc:  # pragma: no cover - the regex already guarantees a valid literal
            raise ValueError("not a number") from exc
    raise ValueError(f"unexpected cell type {type(value).__name__}")


def parse_int_cell(value: object) -> int | None:
    """An integer that may be written as ``"296"`` or ``"296,0000"``. Fractions are rejected."""
    dec = parse_decimal_cell(value)
    if dec is None:
        return None
    if dec != dec.to_integral_value():
        raise ValueError("not a whole number")
    return int(dec)


def parse_text_cell(value: object) -> str | None:
    """Text cells only: a number here (e.g. an EAN stored as a number) would have lost leading zeros."""
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError(f"expected text, got a {type(value).__name__} cell")
    text = value.strip()
    return text or None


def _echo(value: object) -> str:
    text = repr(value)
    return text if len(text) <= _MAX_CELL_ECHO else text[: _MAX_CELL_ECHO - 3] + "..."


def _header_map(header_row: tuple[object, ...]) -> tuple[dict[str, int], list[str]]:
    """Header text -> column index, plus drift notes. Raises GvhSchemaError on missing/duplicate headers."""
    positions: dict[str, int] = {}
    drift: list[str] = []
    for index, cell in enumerate(header_row):
        if cell is None or (isinstance(cell, str) and not cell.strip()):
            continue
        name = _norm(str(cell))
        if name in positions:
            raise GvhSchemaError(f"daily XLSX: duplicated header {name!r}")
        positions[name] = index
        if name not in HEADERS:
            drift.append(f"xlsx.header:{name}")
    missing = [name for name in HEADERS if name not in positions]
    if missing:
        raise GvhSchemaError(f"daily XLSX: missing header(s) {missing!r}")
    if [positions[name] for name in HEADERS] != sorted(positions[name] for name in HEADERS):
        drift.append("xlsx.column_order")
    return {name: positions[name] for name in HEADERS}, drift


def _cell(row: tuple[object, ...], index: int) -> object:
    return row[index] if index < len(row) else None


def _parse_row(row: tuple[object, ...], cols: dict[str, int], row_number: int) -> XlsxRow:
    def get(name: str) -> object:
        return _cell(row, cols[name])

    def fail(name: str, reason: str) -> GvhSchemaError:
        return GvhSchemaError(f"daily XLSX row {row_number}, column {name!r}: {reason} ({_echo(get(name))})")

    def text(name: str) -> str | None:
        try:
            return parse_text_cell(get(name))
        except ValueError as exc:
            raise fail(name, str(exc)) from None

    def required_text(name: str) -> str:
        value = text(name)
        if value is None:
            raise fail(name, "empty")
        return value

    def number(name: str) -> Decimal | None:
        try:
            return parse_decimal_cell(get(name))
        except ValueError as exc:
            raise fail(name, str(exc)) from None

    def whole(name: str) -> int | None:
        try:
            return parse_int_cell(get(name))
        except ValueError as exc:
            raise fail(name, str(exc)) from None

    category_id = whole("Kategória azonosító")
    if category_id is None:
        raise fail("Kategória azonosító", "empty")
    return XlsxRow(
        product_id=required_text("Termék azonosító"),
        product_name=required_text("Termék név"),
        category_id=category_id,
        category_name=text("Kategória név"),
        chain_name=required_text("Üzletlánc név"),
        unit=required_text("Egység"),
        packaging=number("Kiszerelés"),
        min_price=number("Minimum ár"),
        max_price=number("Maximum ár"),
        min_unit_price=number("Minimum egységár"),
        max_unit_price=number("Maximum egységár"),
        chain_count=whole("Hány üzletláncban elérhető"),
        stores_carrying=whole("Hány boltban elérhető"),
        chain_store_count=whole("Üzletlánc összes boltja"),
    )


def _is_blank(row: tuple[object, ...]) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row)


def _read_sheet(workbook: openpyxl.Workbook, expected_date: date) -> DailyXlsx:
    dated = [(name, day) for name in workbook.sheetnames if (day := sheet_date(name)) is not None]
    if len(dated) != 1:
        shown = [_echo(name) for name in workbook.sheetnames[:5]]
        raise GvhSchemaError(
            f"daily XLSX: expected exactly one sheet named '{SHEET_PREFIX}YYYY-MM-DD', found {shown}"
        )
    name, day = dated[0]
    if day != expected_date:
        raise StaleXlsxError(day, expected_date)
    numbered = enumerate(workbook[name].iter_rows(values_only=True), start=1)
    header = next((row for _, row in numbered if not _is_blank(row)), None)
    if header is None:
        raise GvhSchemaError("daily XLSX: the sheet is empty")
    cols, drift = _header_map(header)
    rows = tuple(_parse_row(row, cols, number) for number, row in numbered if not _is_blank(row))
    if not rows:
        raise GvhSchemaError("daily XLSX: no data rows")
    return DailyXlsx(data_date=day, sheet_name=_norm(name), rows=rows, drift=tuple(drift))


XLSX_MAX_ENTRIES = 200
XLSX_MAX_UNPACKED_BYTES = 256 * 1024 * 1024  # the real file unpacks to about 12.6 MB


def _require_hardened_xml() -> None:
    """The XLSX comes from a CDN, i.e. it is untrusted input: refuse to parse it unless openpyxl uses
    defusedxml (entity-expansion and external-entity protection). defusedxml is a runtime dependency."""
    if not DEFUSEDXML:
        raise GvhError("daily XLSX: openpyxl is not using defusedxml; refusing to parse untrusted XML")


def _check_archive(content: bytes) -> None:
    """Reject zip bombs before openpyxl touches the archive. zipfile never inflates a member beyond its
    declared size, so bounding the declared sizes bounds the work."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
    except Exception as exc:  # untrusted bytes: any failure means "not a readable XLSX"
        raise GvhSchemaError(f"daily XLSX: not a readable XLSX file ({type(exc).__name__})") from None
    if len(infos) > XLSX_MAX_ENTRIES or sum(info.file_size for info in infos) > XLSX_MAX_UNPACKED_BYTES:
        raise GvhSchemaError("daily XLSX: archive too large when unpacked")


def parse_daily_xlsx(content: bytes, expected_date: date) -> DailyXlsx:
    """Parse the daily XLSX; raise StaleXlsxError unless its sheet date equals ``expected_date``.

    Anything openpyxl cannot read (not a zip, corrupt XML, read lazily while iterating) is a GvhSchemaError
    naming only the exception type: the input is untrusted and its content is never echoed."""
    _require_hardened_xml()
    _check_archive(content)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # any parser failure on untrusted bytes means "not a readable XLSX"
        raise GvhSchemaError(f"daily XLSX: not a readable XLSX file ({type(exc).__name__})") from None
    try:
        return _read_sheet(workbook, expected_date)
    except GvhError:
        raise
    except Exception as exc:  # lazily parsed sheet XML can fail mid-iteration
        raise GvhSchemaError(f"daily XLSX: unreadable sheet ({type(exc).__name__})") from None
    finally:
        workbook.close()
