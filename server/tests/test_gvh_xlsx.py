"""The daily XLSX parser (BUILD-BRIEF §4.1 trap 4): sheet-date check, decimal-comma text numbers, header
validation. Workbooks are built in memory with openpyxl (synthetic content)."""

from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from bpf.gvh import XLSX_HEADERS, GvhSchemaError, StaleXlsxError, parse_daily_xlsx
from bpf.gvh.xlsx import parse_decimal_cell, parse_int_cell, parse_text_cell, sheet_date
from tests.fakegvh import build_world, xlsx_bytes, xlsx_rows

DAY = date(2026, 10, 13)
ROW = [
    "5997490112823",
    "Próbakert spagetti 500 g",
    "50",
    "Száraztészta",
    "Spar",
    "kg",
    "0,5000",
    "294,0000",
    "699,0000",
    "588,0000",
    "1398,0000",
    "1",
    "298",
    "375",
]


def workbook(
    rows: Sequence[Sequence[object]],
    *,
    header: Sequence[str] = XLSX_HEADERS,
    sheets: Sequence[str] = ("Napi feltöltések - 2026-10-13",),
    leading_blank: bool = False,
) -> bytes:
    wb = openpyxl.Workbook()
    first = wb.active
    assert first is not None
    for index, name in enumerate(sheets):
        sheet = first if index == 0 else wb.create_sheet()
        sheet.title = name
        if leading_blank:
            sheet.append([None])
        sheet.append(list(header))
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ───────────────────────────── cells ─────────────────────────────


@pytest.mark.parametrize(
    ("cell", "expected"),
    [
        ("1699,0000", Decimal("1699.0000")),
        ("0,5", Decimal("0.5")),
        ("  12 ", Decimal(12)),
        ("-3,25", Decimal("-3.25")),
        ("", None),
        ("   ", None),
        (None, None),
        (5, Decimal(5)),
        (0.285, Decimal("0.285")),
    ],
)
def test_decimal_comma_cells(cell: object, expected: Decimal | None) -> None:
    assert parse_decimal_cell(cell) == expected


def test_decimal_comma_parsing_is_exact() -> None:
    value = parse_decimal_cell("1234,5678")
    assert value == Decimal("1234.5678")
    assert isinstance(value, Decimal)
    assert str(value) == "1234.5678"


@pytest.mark.parametrize(
    "cell", ["1.699,00", "1 699,00", "1 699,00", "1699.00", "1,699,00", "12a", "NaN", "Infinity", True, ","]
)
def test_ambiguous_or_invalid_numbers_are_rejected(cell: object) -> None:
    with pytest.raises(ValueError):  # noqa: PT011 - the parser only raises ValueError
        parse_decimal_cell(cell)


@pytest.mark.parametrize(("cell", "expected"), [("296", 296), ("296,0000", 296), (7, 7), ("", None)])
def test_integer_cells(cell: object, expected: int | None) -> None:
    assert parse_int_cell(cell) == expected


def test_fractional_integer_cell_is_rejected() -> None:
    with pytest.raises(ValueError, match="whole"):
        parse_int_cell("296,5")


def test_text_cells() -> None:
    assert parse_text_cell("  abc ") == "abc"
    assert parse_text_cell("") is None
    with pytest.raises(ValueError, match="expected text"):
        parse_text_cell(5997490112823)  # an EAN stored as a number would lose leading zeros


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("Napi feltöltések - 2026-10-13", DAY),
        ("  Napi   feltöltések - 2026-10-13 ", DAY),
        ("Napi feltöltések - 2026-13-40", None),
        ("Napi feltoltesek - 2026-10-13", None),
        ("Sheet1", None),
    ],
)
def test_sheet_date(name: str, expected: date | None) -> None:
    assert sheet_date(name) == expected


# ───────────────────────────── whole files ─────────────────────────────


def test_parses_rows_with_decimal_commas() -> None:
    sheet = parse_daily_xlsx(workbook([ROW]), DAY)
    assert sheet.data_date == DAY
    assert sheet.drift == ()
    (row,) = sheet.rows
    assert row.product_id == "5997490112823"
    assert row.category_id == 50
    assert row.chain_name == "Spar"
    assert row.unit == "kg"
    assert row.packaging == Decimal("0.5")
    assert (row.min_price, row.max_price) == (Decimal(294), Decimal(699))
    assert (row.min_unit_price, row.max_unit_price) == (Decimal(588), Decimal(1398))
    assert (row.chain_count, row.stores_carrying, row.chain_store_count) == (1, 298, 375)


def test_empty_optional_cells_are_none() -> None:
    row = [*ROW[:6], "", *ROW[7:9], "", "", "", "", ""]
    (parsed,) = parse_daily_xlsx(workbook([row]), DAY).rows
    assert parsed.packaging is None
    assert parsed.min_unit_price is None
    assert parsed.stores_carrying is None


def test_stale_sheet_raises() -> None:
    content = workbook([ROW], sheets=("Napi feltöltések - 2026-10-12",))
    with pytest.raises(StaleXlsxError) as err:
        parse_daily_xlsx(content, DAY)
    assert err.value.sheet_date == date(2026, 10, 12)


def test_future_sheet_is_also_not_the_expected_day() -> None:
    with pytest.raises(StaleXlsxError):
        parse_daily_xlsx(workbook([ROW], sheets=("Napi feltöltések - 2026-10-14",)), DAY)


@pytest.mark.parametrize(
    "sheets",
    [("Munka1",), ("Napi feltöltések - 2026-10-13", "Napi feltöltések - 2026-10-12")],
)
def test_sheet_name_must_identify_one_day(sheets: tuple[str, ...]) -> None:
    with pytest.raises(GvhSchemaError, match="exactly one sheet"):
        parse_daily_xlsx(workbook([ROW], sheets=sheets), DAY)


def test_other_sheets_are_ignored() -> None:
    content = workbook([ROW], sheets=("Napi feltöltések - 2026-10-13", "Info"))
    assert len(parse_daily_xlsx(content, DAY).rows) == 1


def test_missing_header_fails_loudly() -> None:
    header = [h for h in XLSX_HEADERS if h != "Maximum ár"]
    with pytest.raises(GvhSchemaError, match="Maximum ár"):
        parse_daily_xlsx(workbook([ROW[:8] + ROW[9:]], header=header), DAY)


def test_renamed_header_fails_loudly() -> None:
    world = build_world()
    with pytest.raises(GvhSchemaError, match="missing header"):
        parse_daily_xlsx(xlsx_bytes(world, DAY, mode="bad_header"), DAY)


def test_duplicated_header_fails_loudly() -> None:
    with pytest.raises(GvhSchemaError, match="duplicated"):
        parse_daily_xlsx(workbook([[*ROW, "x"]], header=[*XLSX_HEADERS, "Termék név"]), DAY)


def test_extra_column_is_drift() -> None:
    world = build_world()
    sheet = parse_daily_xlsx(xlsx_bytes(world, DAY, mode="extra_column"), DAY)
    assert sheet.drift == ("xlsx.header:Új oszlop",)
    assert len(sheet.rows) == len(xlsx_rows(world))


def test_reordered_columns_are_mapped_by_name_and_reported() -> None:
    world = build_world()
    plain = parse_daily_xlsx(xlsx_bytes(world, DAY), DAY)
    reordered = parse_daily_xlsx(xlsx_bytes(world, DAY, mode="reordered"), DAY)
    assert reordered.drift == ("xlsx.column_order",)
    assert reordered.rows == plain.rows


def test_header_whitespace_and_leading_blank_rows_are_tolerated() -> None:
    header = [f" {h}  " for h in XLSX_HEADERS]
    content = workbook([ROW, [None] * 14, ROW], header=header, leading_blank=True)
    assert len(parse_daily_xlsx(content, DAY).rows) == 2


def test_bad_number_names_row_and_column() -> None:
    bad = [*ROW[:7], "1.699,00", *ROW[8:]]
    with pytest.raises(GvhSchemaError, match=r"row 3, column 'Minimum ár'"):
        parse_daily_xlsx(workbook([ROW, bad]), DAY)


def test_numeric_product_id_is_rejected() -> None:
    with pytest.raises(GvhSchemaError, match="Termék azonosító"):
        parse_daily_xlsx(workbook([[5997490112823, *ROW[1:]]]), DAY)


@pytest.mark.parametrize("column", [0, 1, 2, 4, 5])
def test_required_cells(column: int) -> None:
    row = list(ROW)
    row[column] = ""
    with pytest.raises(GvhSchemaError, match="empty"):
        parse_daily_xlsx(workbook([row]), DAY)


def test_no_data_rows_is_an_error() -> None:
    with pytest.raises(GvhSchemaError, match="no data rows"):
        parse_daily_xlsx(workbook([]), DAY)


def test_empty_sheet_is_an_error() -> None:
    wb = openpyxl.Workbook()
    assert wb.active is not None
    wb.active.title = "Napi feltöltések - 2026-10-13"
    buffer = io.BytesIO()
    wb.save(buffer)
    with pytest.raises(GvhSchemaError, match="empty"):
        parse_daily_xlsx(buffer.getvalue(), DAY)


def test_not_an_xlsx_file() -> None:
    with pytest.raises(GvhSchemaError, match="not a readable XLSX"):
        parse_daily_xlsx(b"<!doctype html><html></html>", DAY)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a workbook")
    with pytest.raises(GvhSchemaError, match="not a readable XLSX"):
        parse_daily_xlsx(buffer.getvalue(), DAY)


def test_corrupt_sheet_xml_is_a_schema_error() -> None:
    source = zipfile.ZipFile(io.BytesIO(workbook([ROW])))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = data[: len(data) // 2]  # truncated XML
            target.writestr(item, data)
    with pytest.raises(GvhSchemaError, match="XLSX"):
        parse_daily_xlsx(buffer.getvalue(), DAY)


def test_the_fake_daily_file_round_trips() -> None:
    world = build_world()
    sheet = parse_daily_xlsx(xlsx_bytes(world, DAY), DAY)
    assert len(sheet.rows) == len(xlsx_rows(world))
    spar = next(r for r in sheet.rows if r.product_id == world.special["spar_outlier"])
    assert (spar.min_price, spar.max_price) == (Decimal(294), Decimal(699))  # single-store minimum (trap 1)
    loyalty = next(r for r in sheet.rows if r.product_id == world.special["penny_loyalty"])
    assert loyalty.min_price == loyalty.max_price == Decimal(899)  # the XLSX excludes LOYALTY prices
