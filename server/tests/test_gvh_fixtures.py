"""Every GVH parser over the fixture files: the SYNTHETIC ones (always) and the real samples in
``fixtures/gvh/raw/`` (when someone has committed them; see docs/QUESTIONS.md, 2026-10-13).

The raw tests are how the shape assumptions in ``bpf/gvh/ASSUMPTIONS.md`` get checked: a schema error means
an assumption is wrong; an unknown field means a field to review (and to declare in ``bpf/gvh/models.py``).
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from bpf.gvh.models import (
    is_shop_prices_not_found,
    parse_categories,
    parse_chain_stores,
    parse_product_detail,
    parse_product_page,
    parse_shop_prices,
    parse_shops,
)
from bpf.gvh.xlsx import parse_daily_xlsx, sheet_date
from tests.fakegvh import RAW_DIR, SYNTHETIC_DIR, build_world, synthetic_fixtures, xlsx_bytes


@dataclass(frozen=True)
class FixtureResult:
    kind: str
    drift: tuple[str, ...]
    detail: object = None


def _meta(path: Path) -> dict[str, object]:
    meta_path = path.with_name(path.name.removesuffix(".json") + ".meta.json")
    return json.loads(meta_path.read_text(encoding="utf-8")) if meta_path.exists() else {}


def parse_fixture(path: Path) -> FixtureResult:
    """Run the parser that matches a fixture's name (names as requested in docs/QUESTIONS.md)."""
    name = path.name.removesuffix(".json")
    body = path.read_bytes()
    status = int(_meta(path).get("status", 200))  # type: ignore[call-overload]
    if name == "categories":
        return FixtureResult("categories", parse_categories(body).drift)
    if name == "chain-stores":
        return FixtureResult("chain-stores", parse_chain_stores(body).drift)
    if name.startswith("shops"):
        return FixtureResult("shops", parse_shops(body).drift)
    if name.startswith("pbc-"):
        parsed = parse_product_page(body)
        return FixtureResult("listing", parsed.drift, parsed.value)
    if name.startswith("product-shop"):
        if is_shop_prices_not_found(status, body):
            return FixtureResult("shop-prices-not-found", ())
        assert 200 <= status < 300, f"{path.name}: HTTP {status} without the NOT-FOUND code"
        parsed_shops = parse_shop_prices(body)
        if parsed_shops is None:
            return FixtureResult("shop-prices-not-found", ())
        return FixtureResult("shop-prices", parsed_shops.drift)
    if name.startswith("product-"):
        return FixtureResult("product", parse_product_detail(body).drift)
    if name.startswith("search-"):
        assert status == 400, f"{path.name}: expected HTTP 400"
        assert b"SEARCH-INPUT-ERROR" in body
        return FixtureResult("search-error", ())
    raise AssertionError(f"no parser for fixture {path.name}: add one to tests/test_gvh_fixtures.py")


def parse_xlsx_fixture(path: Path) -> FixtureResult:
    workbook = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True)
    dates = [d for name in workbook.sheetnames if (d := sheet_date(name)) is not None]
    workbook.close()
    assert len(dates) == 1, f"{path.name}: sheets {workbook.sheetnames}"
    sheet = parse_daily_xlsx(path.read_bytes(), dates[0])
    return FixtureResult("xlsx", sheet.drift, sheet)


def _json_files(directory: Path) -> list[Path]:
    return sorted(p for p in directory.glob("*.json") if not p.name.endswith(".meta.json"))


# ───────────────────────────── synthetic (always present) ─────────────────────────────

SYNTHETIC = _json_files(SYNTHETIC_DIR)


def test_synthetic_fixtures_exist_and_are_labelled() -> None:
    assert len(SYNTHETIC) >= 12
    assert "SYNTHETIC" in (SYNTHETIC_DIR / "README.md").read_text(encoding="utf-8")
    assert all(_meta(path).get("synthetic") is True for path in SYNTHETIC)


@pytest.mark.parametrize("path", SYNTHETIC, ids=lambda p: p.name)
def test_synthetic_fixture_parses_without_drift(path: Path) -> None:
    result = parse_fixture(path)
    assert result.drift == ()


def test_synthetic_fixture_kinds() -> None:
    kinds = {path.name: parse_fixture(path).kind for path in SYNTHETIC}
    assert kinds["product-shop-notfound.json"] == "shop-prices-not-found"
    assert kinds["product-shop-ok.json"] == "shop-prices"
    assert kinds["product-multichain.json"] == "product"
    assert kinds["pbc-18-o0.json"] == "listing"


def test_synthetic_default_page_shows_the_12_row_trap() -> None:
    default = parse_fixture(SYNTHETIC_DIR / "pbc-18-default.json").detail
    full = parse_fixture(SYNTHETIC_DIR / "pbc-18-o0.json").detail
    assert default.count == full.count == 101  # type: ignore[union-attr]
    assert len(default.products) == 12  # type: ignore[union-attr]
    assert len(full.products) == 48  # type: ignore[union-attr]


def test_synthetic_xlsx_sample_parses() -> None:
    result = parse_xlsx_fixture(SYNTHETIC_DIR / "xlsx-sample.xlsx")
    assert result.drift == ()
    assert len(result.detail.rows) == 150  # type: ignore[union-attr]


def test_synthetic_fixtures_match_the_generator() -> None:
    """The committed files must be what ``python -m tests.fakegvh --write-fixtures`` writes."""
    for name, (status, body) in synthetic_fixtures().items():
        path = SYNTHETIC_DIR / name
        assert path.exists(), f"missing {name}: regenerate the fixtures"
        assert json.loads(path.read_bytes()) == json.loads(body), f"{name} is stale: regenerate the fixtures"
        assert _meta(path)["status"] == status
    expected = parse_daily_xlsx(
        xlsx_bytes(build_world(), date(2026, 10, 13), row_limit=150), date(2026, 10, 13)
    )
    committed = parse_xlsx_fixture(SYNTHETIC_DIR / "xlsx-sample.xlsx").detail
    assert committed.rows == expected.rows  # type: ignore[union-attr]


# ───────────────────────────── real samples (when committed) ─────────────────────────────

RAW = _json_files(RAW_DIR) if RAW_DIR.is_dir() else []
RAW_XLSX = sorted(RAW_DIR.glob("*.xlsx")) if RAW_DIR.is_dir() else []
_no_raw = pytest.mark.skipif(not RAW_DIR.is_dir(), reason=f"no real GVH samples in {RAW_DIR} yet")


@_no_raw
@pytest.mark.parametrize("path", RAW, ids=lambda p: p.name)
def test_raw_sample_parses(path: Path) -> None:
    parse_fixture(path)


@_no_raw
@pytest.mark.parametrize("path", RAW, ids=lambda p: p.name)
def test_raw_sample_has_no_unknown_fields(path: Path) -> None:
    """Unknown fields pass at run time (as drift warnings); in real samples, review and declare them."""
    assert parse_fixture(path).drift == ()


@_no_raw
@pytest.mark.parametrize("path", RAW_XLSX, ids=lambda p: p.name)
def test_raw_xlsx_sample_parses_without_drift(path: Path) -> None:
    assert parse_xlsx_fixture(path).drift == ()


@_no_raw
def test_raw_default_page_size() -> None:
    """Records the 12-row default page the client must never rely on (BUILD-BRIEF §4.1, §11)."""
    path = RAW_DIR / "pbc-18-default.json"
    if not path.exists():
        pytest.skip("pbc-18-default.json not committed")
    listing = parse_fixture(path).detail
    if listing.count > 12:  # type: ignore[union-attr]
        assert len(listing.products) == 12  # type: ignore[union-attr]
