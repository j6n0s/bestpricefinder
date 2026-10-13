"""tools/capture_gvh_samples.py against the fake GVH: it must produce exactly the samples test_gvh_fixtures.py
reads, every one of them parseable, politely (rate, User-Agent, no redirects)."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

import httpx
import openpyxl
import pytest

from tests.fakegvh import FakeGvh, Recorded, build_world
from tests.test_gvh_fixtures import _json_files, parse_fixture, parse_xlsx_fixture
from tools import capture_gvh_samples as capture

EXPECTED = {
    "categories",
    "chain-stores",
    "shops-head",
    "pbc-18-o0",
    "pbc-18-last",
    "pbc-18-default",
    "pbc-18-shops",
    "pbc-40-o0",
    "pbc-7-o0",
    "product-multichain",
    "product-shop-ok",
    "product-shop-notfound",
    "search-400",
}


Captured = tuple[Path, FakeGvh, dict[str, str], list[Recorded]]


@pytest.fixture(scope="module")
def captured(tmp_path_factory: pytest.TempPathFactory) -> Captured:
    out = tmp_path_factory.mktemp("raw")
    world = build_world()
    fake = FakeGvh(world, data_date=date(2026, 10, 13))
    special = world.special
    with fake.serve() as srv:
        code = capture.main(
            [
                "--contact",
                "ops@example.com",
                "--out",
                str(out),
                "--no-ksh",
                "--base-url",
                srv.api_base,
                "--xlsx-url",
                srv.xlsx_url,
                "--min-interval",
                "0",
                "--product-a",
                special["tesco_zone"],
                "--product-b",
                special["spar_outlier"],
                "--notfound-product",
                special["spar_outlier"],
            ]
        )
    assert code == 0
    return out, fake, special, list(fake.requests)


def test_every_requested_sample_is_written(captured: Captured) -> None:
    out, _, special, _ = captured
    names = {p.name.removesuffix(".json") for p in _json_files(out)}
    assert names == EXPECTED | {f"product-{special['tesco_zone']}", f"product-{special['spar_outlier']}"}
    assert (out / "xlsx-sample.xlsx").exists()
    for path in _json_files(out):
        meta = json.loads(path.with_name(path.stem + ".meta.json").read_text())
        assert {"url", "status", "fetched_at_utc", "content_type", "headers"} <= meta.keys()
        assert "set-cookie" not in {k.lower() for k in meta["headers"]}


def test_every_sample_parses_without_drift(captured: Captured) -> None:
    out, *_ = captured
    kinds = {}
    for path in _json_files(out):
        result = parse_fixture(path)
        assert result.drift == (), path.name
        kinds[path.stem] = result.kind
    assert kinds["product-shop-ok"] == "shop-prices"
    assert kinds["product-shop-notfound"] == "shop-prices-not-found"
    assert kinds["search-400"] == "search-error"
    assert parse_xlsx_fixture(out / "xlsx-sample.xlsx").drift == ()


def test_multichain_sample_has_at_least_four_chains(captured: Captured) -> None:
    out, *_ = captured
    detail = json.loads((out / "product-multichain.json").read_bytes())
    assert len(detail["chainStores"]) >= 4


def test_shops_head_is_truncated_and_untouched(captured: Captured) -> None:
    out, fake, *_ = captured
    head = json.loads((out / "shops-head.json").read_bytes())
    meta = json.loads((out / "shops-head.meta.json").read_text())
    full = json.loads(fake.respond("/api/shops").body)  # type: ignore[union-attr]
    assert head == full[:40]
    assert meta["full_count"] == len(full)


def test_listing_pages_and_shop_ok_requests(captured: Captured) -> None:
    out, *_ = captured
    first = json.loads((out / "pbc-18-o0.json").read_bytes())
    last_url = json.loads((out / "pbc-18-last.meta.json").read_text())["url"]
    assert last_url.endswith(f"limit=48&offset={((first['count'] - 1) // 48) * 48}")
    shop_url = json.loads((out / "product-shop-ok.meta.json").read_text())["url"]
    shop_ids = shop_url.rsplit("/", 1)[1].split(",")
    assert len(shop_ids) == 5
    assert all(s.startswith("tesco-") for s in shop_ids)


def test_xlsx_sample_keeps_sheet_name_header_and_cell_types(
    captured: Captured,
) -> None:
    out, fake, *_ = captured
    meta = json.loads((out / "xlsx-sample.meta.json").read_text())
    sample = openpyxl.load_workbook(out / "xlsx-sample.xlsx", read_only=True)
    served = fake.respond("/" + meta["url"].split("/", 3)[3])  # the same cache-buster: the same file
    original = openpyxl.load_workbook(io.BytesIO(served.body), read_only=True)  # type: ignore[union-attr]
    assert sample.sheetnames == [meta["sheet_name"]] == original.sheetnames[:1]
    sample_rows = list(sample.worksheets[0].iter_rows(values_only=True))
    original_rows = list(original.worksheets[0].iter_rows(values_only=True))
    assert len(sample_rows) == 1 + meta["kept_data_rows"] == 1 + min(150, len(original_rows) - 1)
    assert sample_rows == original_rows[: len(sample_rows)]  # values and their Python types
    assert meta["original_data_rows"] == len(original_rows) - 1


def test_politeness(captured: Captured) -> None:
    *_, requests = captured
    user_agent = "bestpricefinder/0.1 (+contact: ops@example.com; fixture capture)"
    assert {r.user_agent for r in requests} == {user_agent}
    assert len(requests) <= 45
    assert [r.query.get("v", "") for r in requests if r.path.endswith(".xlsx")][0].endswith("-capture")


def test_requests_are_spaced_by_the_minimum_interval() -> None:
    clock = [100.0]
    slept: list[float] = []

    def sleep(seconds: float) -> None:
        slept.append(seconds)
        clock[0] += seconds

    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={}))
    cap = capture.Capture(
        base="https://x.test",
        xlsx_url="https://x.test/f.xlsx",
        out=Path("unused"),
        user_agent="ua",
        min_interval_s=1.2,
        sleep=sleep,
        monotonic=lambda: clock[0],
        client=httpx.Client(transport=transport),
    )
    cap.get("https://x.test/a")
    clock[0] += 0.5
    cap.get("https://x.test/b")
    cap.get("https://x.test/c")
    assert slept == [pytest.approx(0.7), pytest.approx(1.2)]


def test_redirects_are_not_followed_and_ksh_file_is_kept_verbatim(tmp_path: Path) -> None:
    raw = "Kód;Megnevezés;2026. szeptember\n11111;Kenyér;1 234,5\n".encode("iso-8859-2")

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith(".csv"):
            return httpx.Response(200, content=raw, headers={"content-type": "text/csv"})
        return httpx.Response(302, headers={"location": "https://elsewhere.test/"})

    with httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=False) as client:
        cap = capture.Capture(
            base="https://x.test",
            xlsx_url="https://x.test/f.xlsx",
            out=tmp_path / "raw",
            user_agent="ua",
            min_interval_s=0,
            client=client,
        )
        assert cap.get("https://x.test/moved").status_code == 302
        capture.capture_ksh(cap, "https://ksh.test/stadat_files/ara/hu/ara0044.csv", tmp_path / "ksh")
    assert (tmp_path / "ksh" / "ara0044.csv").read_bytes() == raw
    meta = json.loads((tmp_path / "ksh" / "ara0044.csv.meta.json").read_text())
    assert meta["status"] == 200
    assert meta["licence"].startswith("CC BY 4.0")


@pytest.mark.parametrize(
    "argv",
    [
        ["--contact", "nobody"],
        ["--contact", "ops@example.com", "--min-interval", "0.2"],  # real hosts: never faster than 1 req/s
    ],
)
def test_refuses_impolite_settings(argv: list[str]) -> None:
    with pytest.raises(SystemExit) as raised:
        capture.main(argv)
    assert raised.value.code == 2


@pytest.mark.parametrize("status", [403, 429, 503])
def test_stops_at_the_first_refusal_and_never_retries(tmp_path: Path, status: int) -> None:
    world = build_world()
    fake = FakeGvh(world, data_date=date(2026, 10, 13))
    fake.fail("/api/products-by-category/40", status)
    with fake.serve() as srv:
        code = capture.main(
            [
                "--contact",
                "ops@example.com",
                "--out",
                str(tmp_path),
                "--no-ksh",
                "--base-url",
                srv.api_base,
                "--xlsx-url",
                srv.xlsx_url,
                "--min-interval",
                "0",
            ]
        )
    assert code == 1
    assert fake.requests[-1].path == "/api/products-by-category/40"  # nothing after the refusal
    assert len(fake.paths("/api/products-by-category/40")) == 1  # and no retry
    assert not (tmp_path / "pbc-40-o0.json").exists()
    report = (tmp_path / "capture-report.txt").read_text().splitlines()
    assert json.loads(report[-1])["name"] == "STOPPED"
    assert f"HTTP {status}" in json.loads(report[-1])["reason"]


def test_odd_answers_are_saved_as_received_and_optional_items_are_skipped(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        seen.append(path)
        if path == "/api/shops":
            return httpx.Response(200, text="<html>SPA</html>", headers={"content-type": "text/html"})
        if path.startswith("/api/products-by-category/"):
            listed = {"id": "p1", "pricesOfChainStores": [{}, {}, {}]}
            return httpx.Response(200, json={"products": [listed], "count": 1})
        if "/shop/" in path:
            return httpx.Response(404, json={"message": "PRODUCT-SHOP-PRICES-NOT-FOUND"})
        if path.startswith("/api/product/"):
            return httpx.Response(200, json={"chainStores": [{"availableInShops": ["tesco-1"]}, {}]})
        if path == "/api/search":
            return httpx.Response(400, json={"message": "SEARCH-INPUT-ERROR"})
        return httpx.Response(200, json=[{"uuid": "t", "name": "Tesco"}])

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        cap = capture.Capture(
            base="https://x.test/api",
            xlsx_url="https://x.test/f.xlsx",
            out=tmp_path,
            user_agent="ua",
            min_interval_s=0,
            client=client,
        )
        capture.capture_gvh(
            cap,
            product_a="A",
            product_b="B",
            notfound_product="A",
            notfound_store="tesco-4062",
            shop_ids="tesco-4030,lidl-274",
        )
    assert (tmp_path / "shops-head.json").read_text() == "<html>SPA</html>"
    assert json.loads((tmp_path / "shops-head.meta.json").read_text())["full_count"] is None
    assert not (tmp_path / "product-multichain.json").exists()
    assert not (tmp_path / "product-shop-ok.json").exists()
    assert not (tmp_path / "pbc-18-last.json").exists()  # one page only
    assert seen.count("/api/product/p1") == 1  # each candidate is fetched once
    assert seen.count("/api/product/A") == 1  # the named products are not fetched again
    out = capsys.readouterr().out
    assert "none of 1 products tried is sold by >= 4 chains (skipped)" in out
    assert "product-shop-ok           no Budapest Tesco store" in out
